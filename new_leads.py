#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Новые юрлица региона за период → контакты → только те, у кого нет сайта → Excel.

Откуда что берётся:
  1. Список компаний — ЕГРЮЛ (egrul.nalog.ru), перебором ОГРН. С 2022 года
     в регионе один регистрационный центр, номера записей идут подряд, так
     что начало периода находится бинарным поиском по дате регистрации.
     Rusprofile, list-org и прочие справочники скриптам отдают 403 — ФНС
     надёжнее и полнее.
  2. Выписка ЕГРЮЛ — ОКВЭД и адрес (для фильтра по нишам и сверки в 2ГИС).
  3. Контакты — Checko (телефон, почта, сайт из открытых источников),
     по желанию 2ГИС для тех, у кого Checko пуст.
  4. «Сайта нет» — если его нет ни в Checko, ни в 2ГИС, и почта не на
     собственном домене с живым сайтом.

    python new_leads.py --region 25 --since 01.09.2026 --until 30.09.2026
    python new_leads.py --region 25 --since 01.09.2026 --gis vladivostok --all
    python new_leads.py --region 25 --daily --sheet <id Google-таблицы>

--daily — режим для ежедневного запуска по расписанию: берёт последние
60 дней, перепроверяет в Checko тех, у кого контакта ещё не было (Checko
узнаёт о новых компаниях с опозданием), и дописывает в Google-таблицу
только тех, кого там ещё нет. Колонки «Статус», «Кто пишет», «Заметки»
в таблице скрипт не трогает — они для людей.

Промежуточные ответы кешируются в cache_<регион>.json: повторный запуск
за тот же период не ходит в ФНС и Checko заново.
"""

import argparse, csv, datetime as dt, fcntl, json, os, re, sys, time, urllib.parse, urllib.request

from scrapling.fetchers import Fetcher
import logging
logging.getLogger("scrapling").setLevel(logging.CRITICAL)   # «INFO: Fetched…» на каждый запрос — шум

H = {"User-Agent": "Mozilla/5.0", "X-Requested-With": "XMLHttpRequest"}
CA = "/etc/ssl/cert.pem" if os.path.exists("/etc/ssl/cert.pem") else True

# Ниши, где чек от 80 000 ₽ реален. --all отключает фильтр.
NICHES = {
    "медицина":     ("86.1", "86.2", "86.9"),
    "стройка":      ("41.1", "41.2", "43."),
    "архитектура":  ("71.1",),
    "логистика":    ("49.4", "52.2", "52.1"),
    "рестораны":    ("56.1",),
    "мебель":       ("31.0",),
    "автосервис":   ("45.2",),
    "образование":  ("85.4", "85.3"),
    "недвижимость": ("68.3",),
    "производство": ("10.", "16.", "22.", "23.", "25."),
}
# Формы, которым сайт от студии не продать: НКО, фонды, кооперативы, ТСЖ.
SKIP_FORMS = re.compile(r"^(АНО|НКО|ФОНД|ТСЖ|ТСН|СНТ|ГСК|ПК|МКУ|МБУ|ГБУ|ОГКУ|АССОЦИАЦИЯ|СОЮЗ|РЕЛИГИОЗНАЯ|ПОТРЕБИТЕЛЬСКИЙ)\b", re.I)
FREE_MAIL = ("mail.ru", "bk.ru", "inbox.ru", "list.ru", "internet.ru", "yandex.ru", "ya.ru",
             "yandex.com", "gmail.com", "rambler.ru", "icloud.com", "me.com", "outlook.com",
             "hotmail.com", "yahoo.com", "proton.me", "protonmail.com", "vk.com", "lenta.ru",
             "qq.com", "163.com", "126.com", "sina.com", "naver.com", "abv.bg", "mail.com",
             "gmx.com", "gmx.de", "ukr.net", "tut.by", "mail.kz")
MAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
MAIL_JUNK = ("checko", "2gis", "example", "noreply", "sentry", "wixpress", "bitrix", ".png", ".jpg")
NOT_SITE = ("2gis", "vk.com", "t.me", "telegram", "instagram", "facebook", "youtube", "wa.me",
            "whatsapp", "ok.ru", "yandex", "google", "apple", "otello", "sberbank", "gosuslugi")


def log(*a):
    print(*a, flush=True)


# ── кеш ──────────────────────────────────────────────────────────────────────

class Cache:
    def __init__(self, path):
        self.path = path
        self.d = json.load(open(path)) if os.path.exists(path) else {}

    def get(self, kind, key):
        return self.d.get(kind, {}).get(key)

    def put(self, kind, key, val):
        self.d.setdefault(kind, {})[key] = val

    def save(self):
        tmp = self.path + ".tmp"
        json.dump(self.d, open(tmp, "w"), ensure_ascii=False)
        os.replace(tmp, self.path)


# ── ФНС ──────────────────────────────────────────────────────────────────────

def ogrn(year2, region, n):
    base = f"1{year2:02d}{region:02d}00{n:05d}"
    return base + str(int(base) % 11 % 10)


def fns_json(url, data=None, timeout=30):
    req = urllib.request.Request(url, data=data, headers=H)
    return json.load(urllib.request.urlopen(req, timeout=timeout))


def fns_lookup(q):
    """Строка ЕГРЮЛ по ОГРН или None. «Не найдено» приходит без поля i."""
    body = urllib.parse.urlencode({"vyp3CaptchaToken": "", "page": "", "query": q,
                                   "region": "", "PreventChromeAutocomplete": ""}).encode()
    t = fns_json("https://egrul.nalog.ru/", body)["t"]
    for i in range(5):
        time.sleep(1.0)
        rows = [r for r in fns_json(f"https://egrul.nalog.ru/search-result/{t}").get("rows", []) if r.get("i")]
        if rows or i >= 2:
            return rows[0] if rows else None
    return None


def fns_vypiska(tok):
    import pymupdf
    ts = lambda: int(time.time() * 1000)
    t2 = fns_json(f"https://egrul.nalog.ru/vyp-request/{tok}?r=&_={ts()}").get("t", tok)
    for _ in range(20):
        time.sleep(1.5)
        if fns_json(f"https://egrul.nalog.ru/vyp-status/{t2}?r=&_={ts()}").get("status") == "ready":
            break
    pdf = urllib.request.urlopen(urllib.request.Request(
        f"https://egrul.nalog.ru/vyp-download/{t2}", headers=H), timeout=60).read()
    txt = re.sub(r"\s+", " ", " ".join(p.get_text() for p in pymupdf.open("pdf", pdf)))
    okved = re.search(r"основном виде деятельности.*?Код и наименование вида деятельности\s*(\d{2}(?:\.\d{1,2}){0,2})\s+(.+?)\s+\d+\s+ГРН", txt)
    addr = (re.search(r"Адрес юридического лица.*?(\d{6}, .+?)\s+\d{1,3}\s+ГРН", txt)
            or re.search(r"(\d{6}, (?:[А-ЯЁ\-]+ [А-ЯЁ]+|[А-ЯЁ\-]+), .+?)\s+\d{1,3}\s+ГРН", txt))
    return {"оквэд": okved.group(1) if okved else "",
            "вид деятельности": okved.group(2).strip()[:100] if okved else "",
            "адрес": addr.group(1).strip()[:160] if addr else ""}


class Registry:
    """Обращения к ЕГРЮЛ по номеру записи, с кешем."""

    def __init__(self, region, year2, cache):
        self.region, self.year2, self.cache = region, year2, cache

    def at(self, n):
        o = ogrn(self.year2, self.region, n)
        hit = self.cache.get("fns", o)
        if hit is None:
            for attempt in range(3):
                try:
                    r = fns_lookup(o)
                    break
                except Exception as e:
                    log(f"  {o}: ФНС ответила ошибкой {type(e).__name__}, жду")
                    time.sleep(5 * (attempt + 1))
            else:
                return None
            hit = {} if r is None else {
                "огрн": o, "инн": r.get("i"), "компания": r.get("c") or r.get("n"),
                "зарегистрирована": r.get("r"), "директор": re.sub(r"^[А-ЯЁа-яё ]+:\s*", "", r.get("g") or ""),
                "t": r.get("t")}
            self.cache.put("fns", o, hit)
        return hit or None

    def date_near(self, n, span=8):
        """Первая существующая запись начиная с n (в нумерации бывают дыры)."""
        for k in range(n, n + span):
            r = self.at(k)
            if r:
                return k, parse_date(r["зарегистрирована"])
        return None, None

    def first_on_or_after(self, since):
        """Бинарный поиск номера первой записи с датой ≥ since."""
        hi = 64
        while True:                       # удваиваем, пока не перелетим дату или конец года
            k, d = self.date_near(hi)
            if d is None or d >= since:
                break
            hi *= 2
            if hi > 99999:
                break
        lo = 1
        while lo < hi:
            mid = (lo + hi) // 2
            k, d = self.date_near(mid)
            if d is None or d >= since:
                hi = mid
            else:
                lo = (k or mid) + 1
        return lo


def parse_date(s):
    return dt.datetime.strptime(s, "%d.%m.%Y").date()


# ── фильтр ───────────────────────────────────────────────────────────────────

def niche(okved):
    for name, prefixes in NICHES.items():
        if any(okved.startswith(p) for p in prefixes):
            return name
    return ""


def city(addr):
    m = re.search(r"\bГ\. ([А-ЯЁ\-]+)", addr or "")
    return m.group(1).title() if m else ""


# ── контакты ─────────────────────────────────────────────────────────────────

def fetch(url, timeout=30):
    try:
        p = Fetcher.get(url, timeout=timeout, stealthy_headers=True, verify=CA, retries=1)
        return p.status, p.html_content
    except Exception:
        return 0, ""


def plain(html):
    html = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", html, flags=re.S | re.I)
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html))


def field(text, label, stop):
    """Поле блока «Контакты» Checko; ярлык бывает во множественном числе."""
    m = re.search(re.escape(label) + r"[ыа]?\s*(.{0,160}?)\s*(?:" + "|".join(map(re.escape, stop)) + ")", text)
    v = m.group(1).strip() if m else ""
    return "" if v in ("—", "-") else v


CHECKO_KEY = None        # ключ API Checko из checko_key.txt; без него — страница сайта
CHECKO_DAILY = 100       # бесплатный лимит API в сутки
checko_used = 0


class CheckoLimit(Exception):
    pass


def as_list(v):
    if not v:
        return []
    return v if isinstance(v, list) else [v]


def checko_api(o):
    """Контакты через официальный API Checko: без капч и банов, 100 запросов в день бесплатно."""
    global checko_used
    if checko_used >= CHECKO_DAILY - 2:
        raise CheckoLimit()
    url = "https://api.checko.ru/v2/company?" + urllib.parse.urlencode({"key": CHECKO_KEY, "ogrn": o})
    try:
        r = json.load(urllib.request.urlopen(url, timeout=30))
    except Exception as e:
        log(f"  {o}: API Checko — {type(e).__name__}")
        return None
    meta = r.get("meta", {})
    checko_used = meta.get("today_request_count", checko_used + 1)
    if meta.get("status") != "ok":
        log(f"  API Checko: {meta.get('message') or meta}")
        if "лимит" in str(meta).lower() or "limit" in str(meta).lower():
            raise CheckoLimit()
        return None
    k = (r.get("data") or {}).get("Контакты") or {}
    phones = [str(x) for x in as_list(k.get("Тел"))]
    mails = [str(x) for x in as_list(k.get("Емэйл"))]
    sites = [re.sub(r"^https?://(www\.)?", "", str(x)).strip("/") for x in as_list(k.get("ВебСайт"))]
    return {"телефон": "; ".join(phones[:3]), "почта": "; ".join(mails[:2]),
            "сайт": sites[0] if sites else "", "_ts": dt.date.today().isoformat()}


def checko(o):
    """Контакты из Checko: через API, если есть ключ, иначе со страницы компании."""
    if CHECKO_KEY:
        return checko_api(o)
    return checko_page(o)


def checko_page(o):
    """Страница компании на Checko. Сайт режет частые запросы — при 429 ждём и повторяем."""
    for attempt in range(4):
        status, html = fetch(f"https://checko.ru/company/{o}")
        t = plain(html) if status == 200 else ""
        i = t.find("Контакты Адрес")
        if i >= 0:                          # без блока «Контакты» — заглушка, а не «контактов нет»
            block = t[i:i + 700]
            res = {"телефон": field(block, "Телефон", ["Электронная почта", "Электронные"]),
                   "почта": field(block, "Электронная почта", ["Веб-сайт", "Cоциал", "Социал"]),
                   "сайт": field(block, "Веб-сайт", ["Cоциал", "Социал", "Контакты"])}
            if not res["почта"]:
                found = [m for m in MAIL_RE.findall(block) if not any(j in m.lower() for j in MAIL_JUNK)]
                res["почта"] = found[0] if found else ""
            res["_ts"] = dt.date.today().isoformat()
            return res
        time.sleep(10 * (attempt + 1) if status == 429 else 3)
    return None


def gis(row, city_slug):
    """Карточка 2ГИС: ищем по названию, принимаем только если совпала улица или название."""
    name = re.sub(r'^(ООО|АО|ПАО|ЗАО|СЗ)\s*|["«»]', "", row["компания"]).strip()
    m = re.search(r"(?:УЛ\.|ПР-КТ|ПЕР\.|Ш\.|Б-Р)\s*([А-ЯЁ\- ]{3,30}?),", row.get("адрес", ""))
    st = m.group(1).strip().lower() if m else ""
    _, html = fetch(f"https://2gis.ru/{city_slug}/search/{urllib.parse.quote(name)}")
    for fid in list(dict.fromkeys(re.findall(r"/firm/(\d{10,20})", html)))[:3]:
        _, card = fetch(f"https://2gis.ru/{city_slug}/firm/{fid}")
        if not card:
            continue
        title = re.search(r"<title[^>]*>(.*?)</title>", card, re.S)
        core = re.split(r"[,(]", title.group(1) if title else "")[0].strip().lower()
        if not ((st and st in plain(card).lower()) or name.lower()[:9] in core):
            continue
        phones = list(dict.fromkeys("+7" + re.sub(r"\D", "", p)[-10:] for p in re.findall(r"tel:(\+?\d{10,16})", card)))
        sites = [h for h in dict.fromkeys(re.sub(r"^https?://(www\.)?", "", u).split("/")[0].lower()
                                           for u in re.findall(r'href="(https?://[^"]+)"', card))
                 if "." in h and not any(b in h for b in NOT_SITE)]
        mails = [x for x in MAIL_RE.findall(card) if not any(j in x.lower() for j in MAIL_JUNK)]
        return {"телефон": "; ".join(phones[:3]), "почта": mails[0] if mails else "",
                "сайт": sites[0] if sites else "", "2гис": f"https://2gis.ru/{city_slug}/firm/{fid}"}
    return None


def site_alive(host):
    status, html = fetch(f"https://{host}", timeout=15)
    if status == 0:
        status, html = fetch(f"http://{host}", timeout=15)
    parked = re.search(r"domain (is )?for sale|домен продается|reg\.ru.*парк|parking", html or "", re.I)
    return 200 <= status < 400 and len(html) > 1500 and not parked


def site_verdict(r):
    """Возвращает (сайт, пометка)."""
    if r.get("сайт"):
        return r["сайт"], "сайт есть"
    mail = (r.get("почта") or "").split(";")[0].strip().lower()
    dom = mail.split("@")[-1] if "@" in mail else ""
    if dom and dom not in FREE_MAIL and site_alive(dom):
        return dom, "сайт есть (по домену почты)"
    return "", "сайта нет"


# ── выгрузка ─────────────────────────────────────────────────────────────────

COLS = ["компания", "инн", "огрн", "зарегистрирована", "директор", "ниша", "оквэд",
        "вид деятельности", "город", "адрес", "телефон", "почта", "сайт", "пометка",
        "источник", "checko", "2гис"]


def save_xlsx(path, leads, everyone, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    for ws, rows in ((wb.active, leads), (wb.create_sheet(), everyone)):
        ws.title = "Без сайта" if rows is leads else "Все найденные"
        ws.append([c.capitalize() for c in COLS])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2F5496")
        for r in rows:
            ws.append([r.get(c, "") for c in COLS])
        for i, c in enumerate(COLS, 1):
            width = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 50)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("Параметры")
    for k, v in meta.items():
        info.append([k, v])
    wb.save(path)


def save_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:   # BOM — чтобы Excel открыл кириллицу
        w = csv.DictWriter(f, fieldnames=COLS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def has_contact(c):
    return any(c.get(k) for k in ("телефон", "почта", "сайт"))


# ── Google Sheets ────────────────────────────────────────────────────────────

SHEET_COLS = [("Добавлено", None), ("Компания", "компания"), ("ИНН", "инн"), ("ОГРН", "огрн"),
              ("Зарегистрирована", "зарегистрирована"), ("Директор", "директор"), ("Ниша", "ниша"),
              ("Вид деятельности", "вид деятельности"), ("Город", "город"), ("Адрес", "адрес"),
              ("Телефон", "телефон"), ("Почта", "почта"), ("Checko", "checko"), ("2ГИС", "2гис"),
              ("Статус", None), ("Кто пишет", None), ("Заметки", None)]
STATUSES = ["новый", "написали", "ответил", "отказ", "клиент"]


def push_sheet(key_file, sheet_id, leads, meta):
    """Дописывает в лист «Без сайта» тех, чьего ОГРН там ещё нет. Возвращает число добавленных."""
    import gspread
    sh = gspread.service_account(filename=key_file).open_by_key(sheet_id)
    titles = [w.title for w in sh.worksheets()]
    if "Без сайта" in titles:
        ws = sh.worksheet("Без сайта")
    else:
        ws = sh.sheet1 if titles == ["Лист1"] and not any(any(r) for r in sh.sheet1.get_all_values()) else sh.add_worksheet("Без сайта", 1000, 20)
        ws.update_title("Без сайта")
    header = [h for h, _ in SHEET_COLS]
    if ws.row_values(1) != header:
        ws.update([header], "A1", value_input_option="RAW")
        ws.format("A1:Q1", {"textFormat": {"bold": True, "foregroundColor": {"red": 1, "green": 1, "blue": 1}},
                            "backgroundColor": {"red": 0.18, "green": 0.33, "blue": 0.59}})
        ws.freeze(rows=1)
        status_col = header.index("Статус")
        sh.batch_update({"requests": [{"setDataValidation": {
            "range": {"sheetId": ws.id, "startRowIndex": 1, "startColumnIndex": status_col, "endColumnIndex": status_col + 1},
            "rule": {"condition": {"type": "ONE_OF_LIST", "values": [{"userEnteredValue": v} for v in STATUSES]},
                     "showCustomUi": True, "strict": False}}}]})
    known = set(ws.col_values(header.index("ОГРН") + 1)[1:])
    today = f"{dt.date.today():%d.%m.%Y}"
    new = [[today if key is None and h == "Добавлено" else ("новый" if h == "Статус" else (r.get(key, "") if key else ""))
            for h, key in SHEET_COLS]
           for r in leads if r["огрн"] not in known]
    if new:
        ws.append_rows(new, value_input_option="RAW", table_range="A1")
    try:
        log_ws = sh.worksheet("Журнал")
    except gspread.WorksheetNotFound:
        log_ws = sh.add_worksheet("Журнал", 1000, 10)
        log_ws.update([["Запуск", "Период", "Зарегистрировано", "В нишах", "С сайтом",
                        "Без сайта с контактом", "Без контакта", "Не проверено", "Добавлено в таблицу"]], "A1")
        log_ws.freeze(rows=1)
    log_ws.append_row([meta["создано"], meta["период"], meta["зарегистрировано"], meta["в нишах"], meta["с сайтом"],
                       meta["без сайта и с контактом"], meta["без сайта и без контакта"],
                       meta["не проверено (перезапустить позже)"], len(new)], value_input_option="RAW")
    return len(new)


# ── этапы ────────────────────────────────────────────────────────────────────

def collect(region, since, until, cache, limit=None):
    """Компании, зарегистрированные в [since, until]. Номера ОГРН идут заново каждый год,
    поэтому период через Новый год разбивается на куски по годам."""
    companies = []
    for year in range(since.year, until.year + 1):
        a, b = max(since, dt.date(year, 1, 1)), min(until, dt.date(year, 12, 31))
        reg = Registry(region, year % 100, cache)
        log(f"Ищу в ЕГРЮЛ первую запись от {a:%d.%m.%Y} (регион {region})…")
        n = reg.first_on_or_after(a)
        cache.save()
        log(f"  начинаю с номера {n}")
        misses = 0
        while misses < 40:
            r = reg.at(n)
            n += 1
            if n % 20 == 0:
                cache.save()
            if not r:
                misses += 1
                continue
            misses = 0
            d = parse_date(r["зарегистрирована"])
            if d > b:
                break
            if d >= a:
                companies.append(dict(r))
            if limit and len(companies) >= limit:
                break
        cache.save()
    log(f"Зарегистрировано за период: {len(companies)}\n")
    return companies


def enrich(companies, cache, all_niches):
    """Выписки ФНС и отбор по формам и нишам."""
    targets = []
    for i, r in enumerate(companies, 1):
        if SKIP_FORMS.search(r["компания"] or ""):
            continue
        v = cache.get("vyp", r["огрн"])
        if v is None:
            try:
                try:
                    v = fns_vypiska(r["t"])
                except Exception:                  # токен из кеша мог протухнуть — берём свежий
                    v = fns_vypiska(fns_lookup(r["огрн"])["t"])
                cache.put("vyp", r["огрн"], v)
            except Exception as e:
                log(f"  {r['огрн']}: выписка не взялась ({type(e).__name__})")
                v = {"оквэд": "", "вид деятельности": "", "адрес": ""}
        r.update(v)
        r["ниша"] = niche(r["оквэд"])
        r["город"] = city(r["адрес"])
        if all_niches or r["ниша"]:
            targets.append(r)
        if i % 10 == 0:
            cache.save()
            log(f"  выписки: {i}/{len(companies)}, в нишах {len(targets)}")
    cache.save()
    log(f"После фильтра форм и ниш: {len(targets)}\n")
    return targets


def stale(c, days=7):
    """Пустой ответ Checko старше недели — стоит спросить ещё раз."""
    if has_contact(c):
        return False
    ts = c.get("_ts")
    return not ts or (dt.date.today() - dt.date.fromisoformat(ts)).days >= days


def contacts(targets, cache, gis_city=None, recheck=0):
    """Контакты из Checko (+ 2ГИС) и вердикт по сайту. recheck — сколько пустых
    ответов старше недели перепросить за прогон."""
    fails, cooled, stop = 0, False, False
    rechecks = 0
    for i, r in enumerate(targets, 1):
        c = cache.get("checko", r["огрн"])
        want = c is None or (recheck and rechecks < recheck and stale(c))
        if want and not stop:
            if fails >= 3:                    # Checko забанил IP — один раз ждём, потом сдаёмся до следующего прогона
                if cooled:
                    log("  Checko так и не отвечает — остальных проверю в следующий раз")
                    stop = True
                else:
                    log("  Checko отвечает 429 подряд, пауза 10 минут…")
                    time.sleep(600)
                    fails, cooled = 0, True
            if not stop:
                if c is not None:
                    rechecks += 1
                try:
                    fresh = checko(r["огрн"])
                except CheckoLimit:
                    log(f"  Дневной лимит API Checko ({CHECKO_DAILY}) исчерпан — остальных проверю завтра")
                    stop = True
                    fresh = None
                if fresh is not None:
                    c = fresh
                    cache.put("checko", r["огрн"], c)
                    fails = 0
                elif not stop:
                    fails += 1
                time.sleep(0.5 if CHECKO_KEY else 2.5)
        checked = c is not None
        c = c or {}
        r.update({k: c.get(k, "") for k in ("телефон", "почта", "сайт")})
        r["источник"] = "checko" if has_contact(c) else ""
        r["checko"] = f"https://checko.ru/company/{r['огрн']}"
        r["2гис"] = ""
        if gis_city and not r["сайт"] and not r["телефон"]:
            g = cache.get("gis", r["огрн"])
            if g is None:
                g = gis(r, gis_city) or {}
                cache.put("gis", r["огрн"], g)
            for k in ("телефон", "почта", "сайт"):
                r[k] = r[k] or g.get(k, "")
            r["2гис"] = g.get("2гис", "")
            if has_contact(g):
                r["источник"] = (r["источник"] + " + 2гис").strip(" +")
        r["сайт"], r["пометка"] = site_verdict(r)
        if not checked and not r["сайт"] and not r["телефон"]:
            r["пометка"] = "не проверено: Checko не ответил"
        log("%-36s %-18s %-28s %s" % ((r["компания"] or "")[:36], r["телефон"][:18] or "—",
                                      r["почта"][:28] or "—", r["пометка"] if not r["сайт"] else r["сайт"]))
        if i % 10 == 0:
            cache.save()
    cache.save()
    if recheck:
        log(f"\nПерепроверено в Checko: {rechecks}")


# ── main ─────────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--region", type=int, required=True, help="код региона: 25 Приморье, 77 Москва, 78 СПб, 38 Иркутск")
    ap.add_argument("--since", help="дата начала, ДД.ММ.ГГГГ")
    ap.add_argument("--until", help="дата конца включительно (по умолчанию сегодня)")
    ap.add_argument("--daily", action="store_true", help="ежедневный режим: последние 60 дней + перепроверка пустых")
    ap.add_argument("--sheet", help="id Google-таблицы, куда дописывать новые компании")
    ap.add_argument("--key", default="google_key.json", help="ключ сервисного аккаунта Google (по умолчанию рядом со скриптом)")
    ap.add_argument("--all", action="store_true", help="не фильтровать по нишам")
    ap.add_argument("--gis", metavar="CITY", help="добирать контакты в 2ГИС, слаг города: vladivostok, moscow, spb")
    ap.add_argument("--limit", type=int, help="остановиться после N компаний (для пробного прогона)")
    ap.add_argument("--out", help="файл .xlsx (по умолчанию ~/Desktop/Work/Новые_без_сайта_<регион>_<период>.xlsx; в --daily не пишется)")
    a = ap.parse_args()

    until = parse_date(a.until) if a.until else dt.date.today()
    if a.daily:
        since = until - dt.timedelta(days=60)
    elif a.since:
        since = parse_date(a.since)
    else:
        sys.exit("Нужен --since или --daily.")
    here = os.path.dirname(os.path.abspath(__file__))
    key = a.key if os.path.isabs(a.key) else os.path.join(here, a.key)

    # Один прогон на регион за раз: иначе два процесса затирают кеш друг друга.
    lock = open(os.path.join(here, f"cache_{a.region:02d}.lock"), "w")
    try:
        fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except BlockingIOError:
        sys.exit(f"По региону {a.region} уже идёт другой прогон — дождитесь его конца.")
    cache = Cache(os.path.join(here, f"cache_{a.region:02d}.json"))
    global CHECKO_KEY
    key_path = os.path.join(here, "checko_key.txt")
    if os.path.exists(key_path):
        CHECKO_KEY = open(key_path).read().strip()
        log("Контакты: API Checko")
    log(f"=== {dt.datetime.now():%d.%m.%Y %H:%M} · регион {a.region} · {since:%d.%m.%Y} – {until:%d.%m.%Y} ===")

    companies = collect(a.region, since, until, cache, a.limit)
    targets = enrich(companies, cache, a.all)
    contacts(targets, cache, a.gis, recheck=(40 if CHECKO_KEY else 20) if a.daily else 0)

    leads = [r for r in targets if r["пометка"] == "сайта нет" and (r["телефон"] or r["почта"])]
    meta = {"регион": a.region, "период": f"{since:%d.%m.%Y} – {until:%d.%m.%Y}",
            "зарегистрировано": len(companies), "в нишах": len(targets),
            "с сайтом": sum(1 for r in targets if r["сайт"]),
            "без сайта и с контактом": len(leads),
            "без сайта и без контакта": sum(1 for r in targets if r["пометка"] == "сайта нет" and not (r["телефон"] or r["почта"])),
            "не проверено (перезапустить позже)": sum(1 for r in targets if r["пометка"].startswith("не проверено")),
            "ниши": "все" if a.all else ", ".join(NICHES), "создано": f"{dt.datetime.now():%d.%m.%Y %H:%M}"}
    log("\n" + "\n".join(f"{k}: {v}" for k, v in meta.items()))

    if a.sheet:
        added = push_sheet(key, a.sheet, leads, meta)
        log(f"\nВ Google-таблицу добавлено новых: {added}")
    if a.out or not a.daily:
        tag = f"{a.region:02d}_{since:%d.%m}-{until:%d.%m.%Y}"
        out = a.out or os.path.expanduser(f"~/Desktop/Work/Новые_без_сайта_{tag}.xlsx")
        save_xlsx(out, leads, targets, meta)
        save_csv(out[:-5] + ".csv", leads)
        log(f"Файл: {out}")


if __name__ == "__main__":
    main()
